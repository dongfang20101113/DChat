"""自动找 PPT 里的版式事故：形状重叠、超出页面。

为什么需要：这个 deck 已经出过两次同类问题 ——
  1. 注册页规则表 7 行溢出到幻灯片外，最后一行被切
  2. 第 8 页时间线表格高 5.0（到 6.3），卡片从 5.55 开始，盖住表格底部
  
每次都是"渲染出来看一眼才发现"，而 19 页逐页看很慢、也容易漏。
这个脚本直接算几何：谁和谁重叠、谁越界，一遍扫完。

判定说明（避免误报）：
  - 卡片之间的**容器重叠**往往是故意的（比如文字框落在色块卡片上），
    所以只报"重叠面积占较小形状面积的比例"超过阈值的；
  - 相邻卡片共用边框线不算重叠（留 0.03 英寸容差）；
  - 幻灯片可用区域按 13.333 x 7.5 英寸（16:9）算。
"""
import sys
import zipfile

from pptx import Presentation
from pptx.util import Emu

EMU_PER_INCH = 914400
SLIDE_W = 13.333
SLIDE_H = 7.5
TOLERANCE = 0.03          # 英寸，容差（共用边框不算重叠）
OVERLAP_RATIO = 0.18      # 重叠面积超过较小形状的 18% 才报


def inches(value):
    return float(value) / EMU_PER_INCH


def bounds(shape):
    return (inches(shape.left), inches(shape.top),
            inches(shape.left + shape.width), inches(shape.top + shape.height))


def text_height_inches(shape):
    """按文字行数估算**实际占用**的高度（英寸）。

    为什么不能直接用 shape.height：文本框的 frame 往往比文字高得多
    （留白、autofit 前的原尺寸），拿 frame 算重叠会把"文字离得挺远"也报成重叠 ——
    第一版就是这么在封面上误报 4 处的。
    """
    if not shape.has_text_frame:
        return None
    frame = shape.text_frame
    total = 0
    has_text = False
    for paragraph in frame.paragraphs:
        text = "".join(run.text for run in paragraph.runs)
        if not text.strip():
            continue
        has_text = True
        # 字号：取该段第一个 run 的，拿不到就按 18pt 估
        size = 18.0
        for run in paragraph.runs:
            if run.font.size is not None:
                size = float(run.font.size) / 12700.0  # EMU -> pt
                break
        # 行距按 1.35 倍估；10 个中文字符在 13.33 英寸页宽里大约占 10*size/72 英寸宽，
        # 这里只关心高度，所以按"每段至少一行"算，超长段落再折行估算
        width_inch = max(0.1, inches(shape.width))
        # 中文字符宽约等于字号，西文约 0.5 倍。粗略按 0.75 倍混合估算
        chars_per_line = max(1.0, width_inch * 72.0 / (size * 0.75))
        lines = max(1, int(len(text) / chars_per_line) + (1 if len(text) % chars_per_line else 0))
        total += lines * size * 1.35 / 72.0
    if not has_text:
        return None
    return total


def real_bounds(shape):
    """重叠判定用的边界：文本框按实际文字高度收紧，其它形状用原 frame。"""
    left, top, right, bottom = bounds(shape)
    # 表格没有 text_frame（是 graphicFrame），用原框
    height = text_height_inches(shape)
    if height is not None:
        bottom = min(bottom, top + height + 0.06)
    return (left, top, right, bottom)


def overlap_area(a, b):
    dx = min(a[2], b[2]) - max(a[0], b[0])
    dy = min(a[3], b[3]) - max(a[1], b[1])
    if dx <= TOLERANCE or dy <= TOLERANCE:
        return 0.0
    return dx * dy


def area(a):
    return max(0.0, a[2] - a[0]) * max(0.0, a[3] - a[1])


def describe(shape):
    if shape.has_text_frame and shape.text_frame.text.strip():
        text = shape.text_frame.text.strip().replace("\n", " ")[:26]
        return "%s 「%s」" % (shape.shape_type, text)
    return str(shape.shape_type)


def main():
    deck = sys.argv[1] if len(sys.argv) > 1 else r"D:\codes\dchat\dchat-项目介绍.pptx"
    prs = Presentation(deck)
    print("文件:", deck.split("\\")[-1])
    print("页数:", len(prs.slides))
    print()

    total_issues = 0
    for index, slide in enumerate(prs.slides, 1):
        shapes = [s for s in slide.shapes
                  if s.left is not None and s.top is not None
                  and s.width is not None and s.height is not None]
        issues = []

        # ---- 越界 ----
        for shape in shapes:
            box = bounds(shape)
            if box[2] > SLIDE_W + TOLERANCE or box[3] > SLIDE_H + TOLERANCE:
                issues.append("越界: %s  右=%.2f 下=%.2f（页面 %.2f x %.2f）"
                              % (describe(shape), box[2], box[3], SLIDE_W, SLIDE_H))
            if box[0] < -TOLERANCE or box[1] < -TOLERANCE:
                issues.append("越界(左上): %s  左=%.2f 上=%.2f" % (describe(shape), box[0], box[1]))

        # ---- 重叠 ----
        # 先过滤掉"容器 + 内容"这种**故意**的重叠：色块卡片本来就是
        # 一个 AUTO_SHAPE 打底、上面叠若干 TEXT_BOX。第一版没过滤，
        # 19 页报了 83 处，全是这类误报，等于没用。
        def is_container(outer, inner):
            """outer 是否只是 inner 的底板（inner 落在它里面）。"""
            if outer.shape_type != 1:  # 1 = MSO_SHAPE_TYPE.AUTO_SHAPE
                return False
            oa, ob = bounds(outer), bounds(inner)
            return (oa[0] <= ob[0] + TOLERANCE and oa[1] <= ob[1] + TOLERANCE
                    and oa[2] >= ob[2] - TOLERANCE and oa[3] >= ob[3] - TOLERANCE)

        for i in range(len(shapes)):
            for j in range(i + 1, len(shapes)):
                if is_container(shapes[i], shapes[j]) or is_container(shapes[j], shapes[i]):
                    continue
                a, b = real_bounds(shapes[i]), real_bounds(shapes[j])
                shared = overlap_area(a, b)
                if shared <= 0:
                    continue
                smaller = min(area(a), area(b))
                if smaller <= 0:
                    continue
                ratio = shared / smaller
                if ratio >= OVERLAP_RATIO:
                    issues.append("重叠 %.0f%%: %s  <->  %s"
                                  % (ratio * 100, describe(shapes[i]), describe(shapes[j])))

        if issues:
            total_issues += len(issues)
            print("第 %d 页 —— %d 处：" % (index, len(issues)))
            for item in issues:
                print("   ", item)

    print()
    if total_issues:
        print("共 %d 处需要人工确认（有些重叠可能是有意的，但值得逐条看）" % total_issues)
        return 1
    print("✅ 没有发现重叠或越界")
    return 0


if __name__ == "__main__":
    sys.exit(main())
