# 直接读 PPT 文本层核对关键词（比 check_office.py 的 contains 更直观）。
#
# 注意：**图片里的文字读不到**。架构图是 PIL 画的位图，"macOS 客户端"那几个字
# 只存在于像素里，文本层查不到 —— 所以这类内容只能靠渲染图片核对。
import io
import os
import sys

from pptx import Presentation

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DECK = os.path.join(REPO, "dchat-项目介绍.pptx")

# 必须出现在文本层的（四端相关）
MUST_HAVE = [
    "四端互通", "四端都有的功能", "四端字节级一致", "全部组合实测通过",
    "2115", "69 次提交", "17,817",
    "macOS", "真机编译", "四端同时在线", "纵深防御", "各端算出来的字节不一样",
]
# 必须**不再出现**的（过时说法）
MUST_NOT_HAVE = [
    "三端互通", "三个客户端共用", "990", "53 次", "2.7 万",
    "13,852", "只做了静态验证，从未在真机运行过",
]


def main():
    if not os.path.exists(DECK):
        print("找不到", DECK)
        return 1
    prs = Presentation(DECK)

    # 把所有形状里的文字拼起来（含表格与文本框）
    chunks = []
    for index, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.has_text_frame:
                chunks.append((index, shape.text_frame.text))
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        chunks.append((index, cell.text))
    full = "\n".join(text for _, text in chunks)

    print("页数:", len(prs.slides))
    print()
    failed = 0

    print("== 应当出现 ==")
    for needle in MUST_HAVE:
        hit = [page for page, text in chunks if needle in text]
        ok = bool(hit)
        if not ok:
            failed += 1
        print("  %s %-32s %s" % ("OK  " if ok else "FAIL", needle,
                                 ("第 " + ",".join(map(str, hit[:6])) + " 页") if hit else "没找到"))

    print()
    print("== 不应再出现（过时说法）==")
    for needle in MUST_NOT_HAVE:
        hit = [page for page, text in chunks if needle in text]
        ok = not hit
        if not ok:
            failed += 1
        print("  %s %-32s %s" % ("OK  " if ok else "FAIL", needle,
                                 "没了" if ok else ("仍在第 " + ",".join(map(str, hit)) + " 页")))

    print()
    print("图片里的文字读不到（架构图的 macOS 客户端框靠渲染核对）")
    print("%d 项不符" % failed)
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
